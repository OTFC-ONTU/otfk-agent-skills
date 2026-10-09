"""Крос-перевірка робіт на списування: попарна схожість текстів студентів однієї роботи.

  python crosscheck.py "<папка роботи, напр. Перевірка…/УІТП/ПР1>" [--ref "<методичка.docx|.txt>"] [--n 8] [--top 15]

  Витягує текст з docx/pdf/xlsx/pptx/py/txt/ipynb кожної підпапки-студента, відкидає n-грами, що є в методичці
  (--ref), і рахує: спільні n-грами (шт.), Jaccard за 5-грамами, приклади спільних фраз. Сортує пари за підозрілістю.
  Також друкує збіг хешів файлів (однаковий файл у двох студентів) і метадані docx/xlsx (автор, lastModifiedBy, дата створення).
  Потрібні: python-docx, openpyxl, pdfplumber (pip install).
"""
import argparse, hashlib, itertools, json, os, re, sys, zipfile

def text_of(path):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".docx":
            import docx; d = docx.Document(path)
            t = [p.text for p in d.paragraphs]
            for tb in d.tables:
                for r in tb.rows: t.append(" | ".join(c.text for c in r.cells))
            return "\n".join(t)
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf: return "\n".join((p.extract_text() or "") for p in pdf.pages)
        if ext in (".xlsx", ".xlsm"):
            import openpyxl; wb = openpyxl.load_workbook(path, data_only=False); out = []
            for ws in wb:
                for row in ws.iter_rows(values_only=True): out += [str(v) for v in row if v is not None]
            return "\n".join(out)
        if ext == ".pptx":
            with zipfile.ZipFile(path) as z:
                return " ".join(re.sub(r"<[^>]+>", " ", z.read(n).decode("utf8", "ignore")) for n in z.namelist() if n.startswith("ppt/slides/slide"))
        if ext in (".py", ".txt", ".md", ".java", ".js", ".cs", ".html", ".css", ".json", ".ipynb"):
            return open(path, encoding="utf-8", errors="ignore").read()
    except Exception as e:
        return f""
    return ""

def meta_of(path):
    if not path.lower().endswith((".docx", ".xlsx", ".pptx")): return None
    try:
        with zipfile.ZipFile(path) as z:
            x = z.read("docProps/core.xml").decode("utf8", "ignore")
            g = lambda tag: (re.search(rf"<[^>]*{tag}[^>]*>([^<]*)<", x) or [None, ""])[1]
            app = z.read("docProps/app.xml").decode("utf8", "ignore") if "docProps/app.xml" in z.namelist() else ""
            a = (re.search(r"<Application>([^<]*)<", app) or [None, ""])[1]
            return dict(creator=g("creator"), lastModifiedBy=g("lastModifiedBy"), created=g("created"), modified=g("modified"), app=a)
    except Exception:
        return None

def words(t): return re.findall(r"[\w%]+", t.lower().replace("’", "").replace("ʼ", "").replace("'", ""))
def sents(t): return [" ".join(words(x)) for x in re.split(r"[.?!\n:;]+", t) if len(words(x)) >= 3]

def strip_ref(t, refs):
    """Прибрати речення, дослівно взяті з методички (питання, заголовки, умова)."""
    return "\n".join(x for x in sents(t) if x not in refs and not any(x in r for r in refs if len(x) > 25))

def grams(ws, n): return {" ".join(ws[i:i+n]) for i in range(len(ws)-n+1)}

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("folder"); ap.add_argument("--ref"); ap.add_argument("--n", type=int, default=8); ap.add_argument("--top", type=int, default=15)
    a = ap.parse_args()
    ref_t = text_of(a.ref) if a.ref else ""
    ref_w = words(ref_t); refs = set(sents(ref_t))
    ref8, ref5 = grams(ref_w, a.n), grams(ref_w, 5)
    studs = {}
    for s in sorted(os.listdir(a.folder)):
        p = os.path.join(a.folder, s)
        if not os.path.isdir(p): continue
        files = [os.path.join(p, f) for f in os.listdir(p) if os.path.isfile(os.path.join(p, f)) and not f.startswith(("Відгук", "_"))]
        txt = strip_ref("\n".join(text_of(f) for f in files), refs); w = words(txt)
        studs[s] = dict(files=files, w=w, g8=grams(w, a.n) - ref8, g5=grams(w, 5) - ref5,
                        hashes={hashlib.md5(open(f, "rb").read()).hexdigest(): os.path.basename(f) for f in files},
                        meta={os.path.basename(f): meta_of(f) for f in files if meta_of(f)})
    print(f"Студентів: {len(studs)}; n={a.n}; еталон: {a.ref or '—'}\n")
    for s, d in studs.items():
        print(f"  {s}: {len(d['w'])} слів; метадані: {json.dumps(d['meta'], ensure_ascii=False)}")
    rows = []
    for x, y in itertools.combinations(studs, 2):
        A, B = studs[x], studs[y]
        common = A["g8"] & B["g8"]; j = len(A["g5"] & B["g5"]) / max(1, len(A["g5"] | B["g5"]))
        same = set(A["hashes"]) & set(B["hashes"])
        rows.append((len(common), j, x, y, sorted(common, key=len, reverse=True)[:3], [A["hashes"][h] for h in same]))
    rows.sort(key=lambda r: (-(len(r[5]) > 0), -r[0], -r[1]))
    print("\nПари (спільні n-грами поза методичкою | Jaccard-5 | однакові файли | приклади):")
    for c, j, x, y, ex, same in rows[:a.top]:
        flag = "!!! " if same or c >= 20 or j >= 0.25 else ("!  " if c >= 5 or j >= 0.12 else "   ")
        print(f"{flag}{x} ↔ {y}: {c} | {j:.2f} | {same or '-'}")
        for e in ex: print(f"        «{e}»")

if __name__ == "__main__":
    main()
